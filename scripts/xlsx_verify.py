#!/usr/bin/env python3
"""P436: the workbooks pasxlsx saves, read by openpyxl and by LibreOffice.

    ./xlsx dump | xlsx_verify.py [--lo]

testdata/xlsx.paslang with the argument dump writes `tpl1 BASE64` and
`tpl2 BASE64`, the two templates after the cells are written. This reads
each with openpyxl (the cells, the sheets and their states, the merged
cells, the validation, the defined names, the date system) and compares
every entry that was not written with the template's, byte for byte; with
--lo it has LibreOffice (headless, the flatpak on this machine) convert the
first sheet to CSV and compare it with the cells. Python has to be the one
of the venv with openpyxl; LibreOffice runs as the user german.
Exit 1 on a difference.
"""
import base64
import csv
import datetime
import io
import os
import subprocess
import sys
import tempfile
import zipfile

here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, here)
try:
    import openpyxl
except ImportError:
    sys.exit("run with the venv's python: /home/german/.cache/paslang-vartmp/venv-xlsx/bin/python3")

# xlsx_make.py runs main() on import, so take its template functions by exec
src = open(os.path.join(here, "xlsx_make.py")).read().replace("\nmain()\n", "\n")
ns = {}
exec(compile(src, "xlsx_make.py", "exec"), ns)

bad = 0


def check(cond, what):
    global bad
    if not cond:
        bad += 1
        print("BAD", what)


def untouched(tpl, out, changed, gone=()):
    a = zipfile.ZipFile(io.BytesIO(tpl))
    b = zipfile.ZipFile(io.BytesIO(out))
    names = [zi.filename for zi in a.infolist() if zi.filename not in gone]
    check([zi.filename for zi in b.infolist()] == names, "the entries and their order: %s" % [zi.filename for zi in b.infolist()])
    for zi in a.infolist():
        if zi.filename in changed or zi.filename in gone:
            continue
        bi = b.getinfo(zi.filename)
        check(a.read(zi.filename) == b.read(zi.filename), "content of " + zi.filename)
        check((zi.compress_type, zi.CRC, zi.date_time, zi.external_attr) == (bi.compress_type, bi.CRC, bi.date_time, bi.external_attr),
              "metadata of " + zi.filename)
    for name in b.namelist():
        check(b.testzip() is None, "zip test")


def lo_csv(blob, sheet_name):
    with tempfile.TemporaryDirectory(dir="/home/german/.cache/paslang-vartmp") as d:
        os.chmod(d, 0o777)
        p = os.path.join(d, "a.xlsx")
        open(p, "wb").write(blob)
        os.chmod(p, 0o644)
        r = subprocess.run(["runuser", "-u", "german", "--", "env", "HOME=/home/german", "flatpak", "run",
                            "--command=/app/libreoffice/program/soffice", "org.libreoffice.LibreOffice", "--headless",
                            "--convert-to", "csv", "--outdir", d, p], capture_output=True, text=True, timeout=300)
        out = os.path.join(d, "a.csv")
        if not os.path.exists(out):
            print(r.stdout, r.stderr)
            return None
        return list(csv.reader(open(out, encoding="utf-8")))


rows = {}
for ln in sys.stdin:
    if ln.strip():
        n, b = ln.split()
        rows[n] = base64.b64decode(b)

# ---- tpl1
tpl = ns["tpl1"]()
out = rows["tpl1"]
untouched(tpl, out, {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml", "xl/worksheets/sheet1.xml",
                     "xl/worksheets/sheet2.xml"}, gone={"xl/calcChain.xml"})
wb = openpyxl.load_workbook(io.BytesIO(out))
check(wb.sheetnames == ["Datos", "Listas", "Oculta"], "sheet names")
check([w.sheet_state for w in wb.worksheets] == ["visible", "hidden", "veryHidden"], "sheet states")
ws = wb["Datos"]
want = {"A1": "Nombre", "B1": "Importe", "C1": "Fecha", "D1": 1, "A2": "007", "B2": "=B3*2", "C2": datetime.datetime(2024, 3, 31),
        "D2": True, "E2": None, "F2": True, "A3": "Ana & Luis", "B3": 10.5, "C3": datetime.datetime(2024, 10, 24),
        "D3": "  con espacios  ", "E3": "nuevo", "G3": "fuera de spans", "A4": "漢字", "B4": 20.25, "A5": "Total general",
        "B5": "=SUM(B2:B4)", "D5": "era una fórmula", "F5": '="x"&A3', "A6": "ya en línea", "A7": "fila nueva",
        "A8": "entre filas",  "A9": " <a & b>\nx ", "C9": -17}
for ref, v in want.items():
    got = ws[ref].value
    check(got == v, "%s: %r != %r" % (ref, got, v))
check(ws["A2"].number_format == "@", "A2 keeps its text style")
check(ws["C3"].number_format == "dd/mm/yyyy", "C3 keeps its date style")
check(ws["D1"].font.b is False and ws["D1"].number_format == "#,##0.00", "D1 style 4")
check([str(m) for m in ws.merged_cells.ranges] == ["A8:C8"], "merged cells")
check(len(ws.data_validations.dataValidation) == 1 and ws.data_validations.dataValidation[0].formula1 == "Listas!$A$1:$A$3", "validation")
check(ws.dimensions == "A1:G9", "dimension " + ws.dimensions)
check(wb["Listas"]["A4"].value == "cuatro", "hidden sheet cell")
check("Total" in wb.defined_names, "defined name")
check(ws.sheet_properties.tabColor.rgb == "FFFF0000", "tab colour")
if "--lo" in sys.argv:
    c = lo_csv(out, "Datos")
    check(c is not None, "LibreOffice opened tpl1")
    if c:
        check(c[0][:3] == ["Nombre", "Importe", "Fecha"], "LO row 1: %r" % c[0])
        check(c[1][0] == "007", "LO keeps 007: %r" % c[1])
        check(c[2][4] == "nuevo" and c[2][6] == "fuera de spans", "LO row 3: %r" % c[2])
        check(c[4][3] == "era una fórmula", "LO row 5: %r" % c[4])
        check(c[8][0].replace("\r", "") == " <a & b>\nx ", "LO row 9: %r" % c[8])
        print("LibreOffice read tpl1: row 2", c[1], "row 5", c[4])

# ---- tpl2
tpl = ns["tpl2"]()
out = rows["tpl2"]
untouched(tpl, out, {"xl/worksheets/hoja.xml"})
wb = openpyxl.load_workbook(io.BytesIO(out))
ws = wb["Lote"]
check(wb.epoch == datetime.datetime(1904, 1, 1), "1904 epoch")
want = {"A1": "Código", "B1": "Fecha de alta", "C1": "Cantidad", "A2": "0042", "B2": datetime.datetime(2024, 10, 24), "C2": 5,
        "A3": "0043", "B3": 1, "B4": "fila vacía antes"}
for ref, v in want.items():
    got = ws[ref].value
    check(got == v, "tpl2 %s: %r != %r" % (ref, got, v))
check(ws["A3"].number_format == "@", "tpl2 A3 text style")
if "--lo" in sys.argv:
    c = lo_csv(out, "Lote")
    check(c is not None, "LibreOffice opened tpl2")
    if c:
        check(c[0] == ["Código", "Fecha de alta", "Cantidad"], "LO tpl2 row 1: %r" % c[0])
        check(c[1][0] == "0042" and c[1][2] == "5", "LO tpl2 row 2: %r" % c[1])
        print("LibreOffice read tpl2:", c)

print("xlsx verify:", "ok" if not bad else "%d bad" % bad)
sys.exit(1 if bad else 0)
